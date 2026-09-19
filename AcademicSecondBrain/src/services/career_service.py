"""Grounded resume generation and taxonomy-based skill gap analysis."""

import json
import os
import re
from pathlib import Path
from typing import Any

from docx import Document
from llama_index.core.schema import QueryBundle

from src.services.skill_service import TAXONOMY, get_skill_graph, merge_taxonomy_match
"""Factual career data assembly and grounded structured resume generation."""

import json
import os
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from llama_index.core.schema import QueryBundle

from src.rag.ingestion.reader import load_documents_from_path
from src.rag.registry.career import get_profile, get_projects, record_career_run, upsert_profile, upsert_project
from src.rag.registry.skills import get_achievements
from src.services.skill_service import get_skill_graph, merge_taxonomy_match, add_evidence


PROJECT_BULLET_PROMPT = """
You are writing resume bullet points for ONE student project.
Use ONLY facts stated in SOURCE CONTENT below. Do not invent metrics, outcomes, team sizes,
or technologies that are not present in it.

Project title: {title}
Tech stack: {tech_stack}
Target role: {target_role}

SOURCE CONTENT (the only allowed factual basis for your bullets):
\"\"\"
{source_content}
\"\"\"

Return JSON only, no markdown, exactly this shape:
{{"bullets": [{{"text": string}}]}}

Rules:
- 2 to 3 bullets, each starting with a strong past-tense action verb.
- Every number, percentage, or specific technology named in a bullet must appear in SOURCE CONTENT.
- No first-person pronouns. No student name or GPA.
"""

CERTIFICATE_EXTRACTION_PROMPT = """
Extract structured data from the certificate text below.
Return JSON only, exactly this shape:
{{"title": string, "issuer": string, "date": string or null, "skills_mentioned": [string, ...]}}
Only include skills explicitly named or clearly implied by the certificate's own title/content.
If unsure, omit.

Certificate text:
\"\"\"
{certificate_text}
\"\"\"
"""

GAP_ANALYSIS_PROMPT = """
Extract only the technical and professional skills explicitly required by the job description below.
Match skills to common names where possible, but do not invent skills that are not present.
Return JSON only, with exactly this shape:
{{"skills": ["Python", "SQL"]}}
Do not include markdown, explanations, or any text outside the JSON object.

Job description:
\"\"\"
{job_description_text}
\"\"\"
"""


with (Path(__file__).with_name("skill_categories.json")).open(encoding="utf-8") as category_file:
    SKILL_CATEGORIES = json.load(category_file)


def get_skill_category(skill_name: str) -> str:
    return SKILL_CATEGORIES.get(skill_name, "Other")


def _parse_json_object(raw_output: str) -> dict[str, Any]:
    candidate = raw_output.strip()
    fenced = re.search(r"```(?:json)?\s*(.*?)\s*```", candidate, re.DOTALL | re.IGNORECASE)
    if fenced:
        candidate = fenced.group(1)
    try:
        value = json.loads(candidate)
    except json.JSONDecodeError as error:
        raise ValueError("Career model returned invalid JSON") from error
    if not isinstance(value, dict):
        raise ValueError("Career model must return a JSON object")
    return value


def group_skills_by_category(graph: dict[str, Any]) -> dict[str, list[str]]:
    grouped: dict[str, list[str]] = {}
    for skill in graph["skills"]:
        if skill.get("skill_type") != "taxonomy":
            continue
        if not any(evidence.get("source_type") != "quiz" for evidence in skill.get("evidence", [])):
            continue
        category = get_skill_category(skill["skill_name"])
        grouped.setdefault(category, []).append(skill["skill_name"])
    return grouped


def _number_tokens(text: str) -> set[str]:
    return set(re.findall(r"\b\d+(?:\.\d+)?%?\b", text))


async def generate_project_bullets(project: dict[str, Any], llm: Any, target_role: str | None = None) -> list[dict[str, str]]:
    source_content = project.get("description_text", "").strip()
    if not source_content:
        return []
    prompt = PROJECT_BULLET_PROMPT.format(
        title=project["title"],
        tech_stack=project["tech_stack"],
        target_role=target_role or "General software engineering",
        source_content=source_content,
    )
    response = await llm.acomplete(prompt)
    payload = _parse_json_object(getattr(response, "text", str(response)))
    bullets = payload.get("bullets")
    if not isinstance(bullets, list):
        raise ValueError("Project response must contain a bullets list")
    source_numbers = _number_tokens(source_content)
    validated = []
    seen = set()
    for bullet in bullets:
        text = bullet.get("text") if isinstance(bullet, dict) else None
        if not isinstance(text, str) or not text.strip():
            raise ValueError("Project bullet must contain non-empty text")
        if not _number_tokens(text).issubset(source_numbers):
            raise ValueError("Project bullet contains an unsupported number")
        if text.strip().casefold() not in seen:
            seen.add(text.strip().casefold())
            validated.append({"text": text.strip()})
    return validated[:3]


async def build_resume_data(student_id: str, llm: Any, target_role: str | None = None) -> dict[str, Any]:
    profile = get_profile(student_id)
    if profile is None:
        raise ValueError("Student profile is required before generating a resume")
    graph = get_skill_graph(student_id)
    projects = []
    for project in get_projects(student_id):
        bullets = await generate_project_bullets(project, llm, target_role)
        projects.append({**project, "bullets": bullets})
    achievements = get_achievements(student_id)
    certifications = [
        achievement for achievement in achievements
        if json.loads(achievement.get("metadata_json", "{}") or "{}").get("source") == "certification"
    ]
    regular_achievements = [achievement for achievement in achievements if achievement not in certifications]
    return {
        "profile": profile,
        "skills_by_category": group_skills_by_category(graph),
        "projects": projects,
        "certifications": certifications,
        "achievements": regular_achievements,
    }


def _set_cell_text(cell: Any, text: str, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def _add_section(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(35, 35, 35)
    paragraph._p.get_or_add_pPr().append(__import__("docx").oxml.OxmlElement("w:pBdr"))


def render_resume_docx(data: dict[str, Any], output_path: str) -> None:
    profile = data["profile"]
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(8)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(0)
    name = heading.add_run(profile["full_name"])
    name.bold = True
    name.font.size = Pt(16)
    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    contact.add_run(" | ".join(filter(None, [profile.get("location"), profile.get("email"), profile.get("phone")]))).font.size = Pt(8)

    _add_section(document, "Education")
    education = document.add_table(rows=1, cols=2)
    education.autofit = True
    left, right = education.rows[0].cells
    _set_cell_text(left, f"{profile.get('college_name', '')}\n{profile.get('degree', '')}\n{profile.get('branch', '')}\nCGPA: {profile.get('cgpa', '')}")
    _set_cell_text(right, f"{profile.get('college_start', '')} - {profile.get('college_end', '')}")
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if profile.get("school_name"):
        row = education.add_row().cells
        _set_cell_text(row[0], f"{profile['school_name']}\n{profile.get('school_detail', '')}\nPercentage: {profile.get('school_score', '')}")
        _set_cell_text(row[1], profile.get("school_dates", ""))
        row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    if data["skills_by_category"]:
        _add_section(document, "Skills")
        for category, skills in data["skills_by_category"].items():
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run(f"{category}: ").bold = True
            paragraph.add_run(", ".join(skills))

    if data["projects"]:
        _add_section(document, "Projects")
        for project in data["projects"]:
            row = document.add_table(rows=1, cols=2).rows[0].cells
            _set_cell_text(row[0], project["title"], bold=True)
            _set_cell_text(row[1], project["tech_stack"], bold=True)
            row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for bullet in project["bullets"]:
                paragraph = document.add_paragraph(bullet["text"], style="List Bullet 2")
                paragraph.paragraph_format.space_after = Pt(0)

    for title, items in (("Certifications", data["certifications"]), ("Achievements", data["achievements"])):
        if items:
            _add_section(document, title)
            for item in items:
                paragraph = document.add_paragraph(item.get("name", ""), style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(0)
    document.save(output_path)


async def generate_resume(student_id: str, llm: Any, target_role: str | None = None, output_dir: str = "outputs") -> dict[str, Any]:
    data = await build_resume_data(student_id, llm, target_role)
    os.makedirs(output_dir, exist_ok=True)
    output_path = Path(output_dir) / f"{student_id}_resume.docx"
    render_resume_docx(data, str(output_path))
    record_career_run(student_id, "resume", {"target_role": target_role, "project_count": len(data["projects"])})
    return {"student_id": student_id, "target_role": target_role, "file_path": str(output_path), "data": data}


async def suggest_skill_gaps(student_id: str, job_description_text: str, llm: Any) -> dict[str, Any]:
    if not job_description_text.strip():
        raise ValueError("job_description_text is required")
    response = await llm.acomplete(
        GAP_ANALYSIS_PROMPT.format(job_description_text=job_description_text.strip())
    )
    payload = _parse_json_object(getattr(response, "text", str(response)))
    raw_skills = payload.get("skills")
    if not isinstance(raw_skills, list) or any(not isinstance(skill, str) for skill in raw_skills):
        raise ValueError("Career model response must contain a skills list")
    required = []
    for raw_skill in raw_skills:
        canonical = merge_taxonomy_match(raw_skill)
        if canonical and canonical not in required:
            required.append(canonical)
    existing = {skill["skill_name"] for skill in get_skill_graph(student_id)["skills"] if skill.get("skill_type") == "taxonomy"}
    result = {"student_id": student_id, "required_skills": required, "gaps": [{"skill_name": skill, "current_confidence": 0.0} for skill in required if skill not in existing]}
    record_career_run(student_id, "gap-analysis", {"required_skills": required, "gaps": result["gaps"]})
    return result


async def extract_certification(
    file_path: str,
    llm: Any,
    student_id: str,
    original_filename: str | None = None,
) -> dict[str, Any]:
    documents = load_documents_from_path([file_path])
    text = "\n".join(document.text for document in documents)
    filename_hint = Path(original_filename or file_path).stem.replace("_", " ").replace("-", " ").strip()
    if not text.strip() and not filename_hint:
        raise ValueError("Could not extract text from certificate")
    response = await llm.acomplete(
        CERTIFICATE_EXTRACTION_PROMPT.format(
            certificate_text=text or f"Certificate filename: {filename_hint}",
        )
    )
    payload = _parse_json_object(getattr(response, "text", str(response)))
    title = payload.get("title")
    if not isinstance(title, str) or not title.strip():
        title = filename_hint
    if not title:
        raise ValueError("Certificate title is required")
    from src.rag.registry.skills import insert_achievement
    achievement_id = insert_achievement(student_id, title, payload.get("issuer"), payload.get("date"), json.dumps({"source": "certification"}))
    unmatched = []
    for skill in payload.get("skills_mentioned", []):
        try:
            add_evidence(student_id, skill, "certification", title, 0.75)
        except ValueError:
            unmatched.append(skill)
    return {"achievement_id": achievement_id, "title": title, "unmatched_skills": unmatched}